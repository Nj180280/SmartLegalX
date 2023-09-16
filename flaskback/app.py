from flask import Flask, request, send_file, jsonify
from docx import Document
from datetime import datetime  # Import the datetime module

from flask_cors import CORS

app = Flask(__name__)
CORS(app)  # Add this line to enable CORS for your Flask app

def format_date(date):
    try:
        # Parse the formatted date and reformat it as needed
        parsed_date = datetime.strptime(date, '%B %d, %Y')
        formatted_date = parsed_date.strftime('%d-name day of %B %Y-number year')
    except ValueError:
        formatted_date = ''  # Handle invalid date format gracefully if needed
    return formatted_date


@app.route('/generate_will', methods=['POST'])
def generate_document():
    try:
        # Get user input from the form (assuming data is sent as JSON)
        user_data = request.json

        # Load the template document
        template_doc = Document('assets/WillDeed.docx')
        
        formatted_date = format_date(user_data.get('date', ''))

        # Replace placeholders in the template with user data
        for paragraph in template_doc.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace('user-name-input', user_data.get('name', ''))
                run.text = run.text.replace('son-of', user_data.get('sonOf', ''))
                run.text = run.text.replace('abc-address-abc-colony', user_data.get('residence', ''))
                run.text = run.text.replace('age-entry', user_data.get('age', ''))
                run.text = run.text.replace('religion-input', user_data.get('religion', ''))
                run.text = run.text.replace('occupation-name', user_data.get('occupation', ''))
                

                run.text = run.text.replace('Will-receiver-name', user_data.get('willReceiverName', ''))
                run.text = run.text.replace('will-receiver-address', user_data.get('willReceiverAddress', ''))
                run.text = run.text.replace('will-receiver-age', user_data.get('willReceiverAge', ''))
                run.text = run.text.replace('will-receiver-religion', user_data.get('willReceiverReligion', ''))
                run.text = run.text.replace('will-receiver-occupation', user_data.get('willReceiverOccupation', ''))
                run.text = run.text.replace('Will-receiver2-name', user_data.get('willReceiver2Name', ''))
                run.text = run.text.replace('will-receiver2-address', user_data.get('willReceiver2Address', ''))
                run.text = run.text.replace('will-receiver2-age', user_data.get('willReceiver2Age', ''))
                run.text = run.text.replace('will-receiver2-religion', user_data.get('willReceiver2Religion', ''))
                run.text = run.text.replace('will-receiver2-occupation', user_data.get('willReceiver2Occupation', ''))
                run.text = run.text.replace('Will-receiver3-name', user_data.get('willReceiver3Name', ''))
                run.text = run.text.replace('will-receiver3-address', user_data.get('willReceiver3Address', ''))
                run.text = run.text.replace('will-receiver3-age', user_data.get('willReceiver3Age', ''))
                run.text = run.text.replace('will-receiver3-religion', user_data.get('willReceiver3Religion', ''))
                run.text = run.text.replace('will-receiver3-occupation', user_data.get('willReceiver3Occupation', ''))
          
                
                run.text = run.text.replace('family-members', user_data.get('familyMembers', ''))
                run.text = run.text.replace('PropertyInput1', user_data.get('property', ''))

                
                run.text = run.text.replace('User-input-name-wife', user_data.get('userWife', ''))
                
                run.text = run.text.replace('Child1-name', user_data.get('child1Name', ''))
                run.text = run.text.replace('Child2-name', user_data.get('child2Name', ''))
                run.text = run.text.replace('Child3-name', user_data.get('child3Name', ''))
                
                run.text = run.text.replace('day-name day of month-name month of year-number year', formatted_date)

                run.text = run.text.replace('witness1-name', user_data.get('witness1Name', ''))
                run.text = run.text.replace('witness2-name', user_data.get('witness2Name', ''))

        # Save the filled-out document
        filled_doc_path = 'filled_legal_document.docx'
        template_doc.save(filled_doc_path)

        # Send the filled document back to the user for download
        return send_file(filled_doc_path, as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)})


if __name__ == '__main__':
    app.run(debug=True)
